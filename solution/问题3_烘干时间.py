"""问题3  烘干时间（固定半径，附录3 变物性，全域含水率事件判终点）

运行：python solution/问题3_烘干时间.py [N]
输出：result/result3.xlsx（工作表 Sheet1，含临界时刻行）
"""
from __future__ import annotations

import sys
from pathlib import Path

import numpy as np
import openpyxl
import scipy.sparse as sp
from scipy.integrate import solve_ivp
from scipy.optimize import brentq

ROOT = Path(__file__).resolve().parents[1]
INPUT_DIR = ROOT / "inputs" if (ROOT / "inputs").exists() else ROOT
TEMPLATE_DIR = (INPUT_DIR / "templates") if (INPUT_DIR / "templates").exists() else (ROOT / "附件3")

R0, T0, C0 = 0.02, 28.0, 2.55
HT, HC = 25.0, 8.0e-7
THRESHOLD = 0.15                     # 烘干判据 max C < 0.15 kg/kg

N = int(sys.argv[1]) if len(sys.argv) > 1 else 320
POS = np.arange(21) * 0.001
LABELS = [round(i * 0.1, 10) for i in range(21)]
T_EARLY, MAX_STEP_EARLY, MAX_STEP_LATE = 14400.0, 60.0, 600.0
T_END = 2.0e6                        # 事件未触发时的积分上限


def chamber():
    """附件1：烘房温度与空气含湿量；4 h 之后取附件末 1 h 均值。"""
    wb = openpyxl.load_workbook(INPUT_DIR / "附件1.xlsx", data_only=True, read_only=True)
    rows = [r for r in wb["Sheet1"].iter_rows(min_row=2, values_only=True) if r[0] is not None]
    wb.close()
    t = np.array([r[0] for r in rows], float)
    ta = np.array([r[1] for r in rows], float)
    ca = np.array([r[2] for r in rows], float)
    tail = t >= t[-1] - 3600.0
    ta_tail, ca_tail = ta[tail].mean(), ca[tail].mean()

    def env(tt):
        scalar = np.ndim(tt) == 0
        tt = np.atleast_1d(np.asarray(tt, float))
        x = np.clip(tt, t[0], t[-1])
        a, b = np.interp(x, t, ta), np.interp(x, t, ca)
        beyond = tt > t[-1]
        a, b = np.where(beyond, ta_tail, a), np.where(beyond, ca_tail, b)
        return (float(a[0]), float(b[0])) if scalar else (a, b)

    return env


def build_grid():
    dr = R0 / N
    r = np.arange(N + 1) * dr
    faces = np.empty(N + 2)
    faces[1:-1] = 0.5 * (r[:-1] + r[1:])
    faces[0], faces[-1] = 0.0, R0
    vol = np.pi * (faces[1:] ** 2 - faces[:-1] ** 2)
    area = 2 * np.pi * faces
    return r, dr, vol, area[:-1], area[1:]


def harmonic(a):
    s = a[1:] + a[:-1]
    out = np.zeros_like(s)
    nz = s != 0
    out[nz] = 2 * a[1:][nz] * a[:-1][nz] / s[nz]
    return out


def rho_of(C):
    return 650.0 + 128.0 * C


def cp_of(C):
    return 1450.0 + 2736.0 * C / (C + 1.0)


def k_of(C):
    return 0.21 + 0.38 * C / (C + 1.0)


def d_of(C, T):
    return 2.4e-3 * np.exp(-0.45 * C) * np.exp(-3850.0 / (T + 273.15))


def make_rhs(env, geom):
    r, dr, vol, al, ar = geom
    n = N + 1

    def fun(t, y):
        T, C = y[:n], y[n:]
        ta, ca = env(t)
        rho, cp, k, d = rho_of(C), cp_of(C), k_of(C), d_of(C, T)
        kh, dh = harmonic(k), harmonic(d)
        qt, qc = np.zeros(n + 1), np.zeros(n + 1)
        qt[1:n] = -kh * (T[1:] - T[:-1]) / dr
        qc[1:n] = -dh * (C[1:] - C[:-1]) / dr
        qt[n] = HT * (T[-1] - ta)
        qc[n] = HC * (C[-1] - ca)
        dt = (al * qt[:-1] - ar * qt[1:]) / (rho * cp * vol)
        dc = (al * qc[:-1] - ar * qc[1:]) / vol
        return np.concatenate((dt, dc))

    return fun


def jac_pattern():
    n = N + 1
    rows, cols = [], []
    for i in range(n):
        for j in (i - 1, i, i + 1):
            if 0 <= j < n:
                rows += [i, i, n + i, n + i]
                cols += [j, n + j, j, n + j]
    return sp.csr_matrix((np.ones(len(rows), bool), (rows, cols)), shape=(2 * n, 2 * n))


class Path:
    """拼接多段积分的稠密输出。"""

    def __init__(self, size):
        self.size = size
        self.legs = []

    def add(self, t0, sol):
        self.legs.append((float(t0), float(sol.t[-1]), sol))

    def __call__(self, t):
        t = np.atleast_1d(np.asarray(t, float))
        out = np.full((self.size, t.size), np.nan)
        for t0, t1, sol in self.legs:
            m = (t >= t0 - 1e-9) & (t <= t1 + 1e-9)
            if m.any():
                out[:, m] = sol.sol(t[m])
        return out


def integrate(fun, y0, atol):
    """分段积分并监测 max C − 0.15 的首次穿越。"""
    event = lambda t, y: float(np.max(y[N + 1:]) - THRESHOLD)
    event.terminal, event.direction = True, -1.0
    path = Path(2 * (N + 1))
    y, t, hit = y0, 0.0, None
    for t0, t1, step in ((0.0, T_EARLY, MAX_STEP_EARLY), (T_EARLY, T_END, MAX_STEP_LATE)):
        sol = solve_ivp(fun, (max(t0, t), t1), y, method="BDF", rtol=1e-8, atol=atol,
                        jac_sparsity=jac_pattern(), dense_output=True,
                        max_step=step, events=event)
        if not sol.success:
            raise RuntimeError(sol.message)
        path.add(max(t0, t), sol)
        y, t = sol.y[:, -1], float(sol.t[-1])
        if sol.t_events[0].size:
            hit = float(sol.t_events[0][0])
            break
    return path, y, t, hit


def refine(path, t_hit):
    """事件后补积一段，用 brentq 在异号区间上复核临界时刻。"""
    t_lo = max(0.0, t_hit - 120.0)
    t_hi = t_hit + 150.0
    g = lambda tt: float(np.max(path(np.array([tt]))[N + 1:]) - THRESHOLD)
    return brentq(g, t_lo, t_hi, xtol=1e-12)


def sample(path, r, times):
    n = N + 1
    y = path(np.asarray(times, float))
    T = np.array([np.interp(POS, r, v[:n]) for v in y.T])
    C = np.array([np.interp(POS, r, v[n:]) for v in y.T])
    return T, C


def export(times, data, out_name):
    wb = openpyxl.load_workbook(TEMPLATE_DIR / "result3.xlsx")
    ws = wb["Sheet1"]
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    ws.cell(1, 1).value = "时间\\到药材中心的距离"
    for j, lab in enumerate(LABELS):
        ws.cell(1, 2 + j).value = lab
    for i, tt in enumerate(times):
        cell = ws.cell(2 + i, 1)
        cell.value = float(tt)
        cell.number_format = "0.####"
        for j in range(len(LABELS)):
            v = data[i, j]
            if v is None or not np.isfinite(v):
                continue
            c = ws.cell(2 + i, 2 + j)
            c.value = float(v)
            c.number_format = "0.0000"
    out = ROOT / "result" / out_name
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    wb.close()
    return out


def main():
    env = chamber()
    geom = build_grid()
    fun = make_rhs(env, geom)
    y0 = np.concatenate((np.full(N + 1, T0), np.full(N + 1, C0)))
    atol = np.concatenate((np.full(N + 1, 1e-8), np.full(N + 1, 1e-10)))

    path, y_end, t_end, t_hit = integrate(fun, y0, atol)
    if t_hit is None:
        raise RuntimeError("在 t=%.0f s 内未触发含水率事件" % T_END)

    # 事件后补积 150 s 以形成异号区间，再独立复核
    extra = solve_ivp(fun, (t_hit, t_hit + 150.0), y_end, method="BDF", rtol=1e-8,
                      atol=atol, jac_sparsity=jac_pattern(), dense_output=True,
                      max_step=MAX_STEP_EARLY)
    path.add(t_hit, extra)
    t_star = refine(path, t_hit)

    times = np.append(np.arange(1, int(np.floor(t_star / 60.0)) + 1) * 60.0, t_star)
    T, C = sample(path, geom[0], times)
    out = export(times, C, "result3.xlsx")

    print("问题3  烘干时间（固定半径）  N=%d" % N)
    print("  临界烘干时间 t* = %.6f s = %.4f h" % (t_star, t_star / 3600.0))
    print("  临界时刻中心含水率 %.6f，表面含水率 %.6f kg/kg" % (C[-1, 0], C[-1, -1]))
    print("  整数安全时刻 %d s：全域最大含水率 %.6f < 0.15" %
          (int(np.ceil(t_star)), np.max(sample(path, geom[0], [np.ceil(t_star)])[1])))
    print("  结果已写入 %s" % out)


if __name__ == "__main__":
    main()
