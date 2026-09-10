"""问题4  收缩模型（附件2 半径 R(t)，材料坐标，附录4 变物性）

运行：python solution/问题4_收缩模型.py [N]
输出：result/result4.xlsx（工作表 Sheet1，域外留空，末列为药材表面）
"""
from __future__ import annotations

import sys
from pathlib import Path

import numpy as np
import openpyxl
import scipy.sparse as sp
from scipy.integrate import solve_ivp
from scipy.interpolate import PchipInterpolator
from scipy.optimize import brentq

ROOT = Path(__file__).resolve().parents[1]
INPUT_DIR = ROOT / "inputs" if (ROOT / "inputs").exists() else ROOT
TEMPLATE_DIR = (INPUT_DIR / "templates") if (INPUT_DIR / "templates").exists() else (ROOT / "附件3")

R0, T0, C0 = 0.02, 28.0, 2.55
HT, HC = 25.0, 8.0e-7
THRESHOLD = 0.15

N = int(sys.argv[1]) if len(sys.argv) > 1 else 320
POS = np.arange(21) * 0.001                      # 固定物理位置 0~2 cm
LABELS = [round(i * 0.1, 10) for i in range(21)] + ["药材表面"]
T_EARLY, MAX_STEP_EARLY, MAX_STEP_LATE = 14400.0, 60.0, 600.0
T_END = 4.0e6


def chamber():
    """附件1：烘房温度与空气含湿量；4 h 之后取附件末 1 h 均值。"""
    wb = openpyxl.load_workbook(INPUT_DIR / "附件1.xlsx", data_only=True, read_only=True)
    rows = [r for r in wb["Sheet1"].iter_rows(min_row=2, values_only=True) if r[0] is not None]
    wb.close()
    t = np.array([r[0] for r in rows], float)
    ta = np.array([r[1] for r in rows], float)
    ca = np.array([r[2] for r in rows], float)
    tail = t >= t[-1] - 3600.0
    ta_tail, ca_tail = ta[tail].mean(), ca[tail].mean()

    def env(tt):
        scalar = np.ndim(tt) == 0
        tt = np.atleast_1d(np.asarray(tt, float))
        x = np.clip(tt, t[0], t[-1])
        a, b = np.interp(x, t, ta), np.interp(x, t, ca)
        beyond = tt > t[-1]
        a, b = np.where(beyond, ta_tail, a), np.where(beyond, ca_tail, b)
        return (float(a[0]), float(b[0])) if scalar else (a, b)

    return env


def radius_function():
    """附件2 半径（cm → m），保形单调插值，超出数据末端保持末值。"""
    wb = openpyxl.load_workbook(INPUT_DIR / "附件2.xlsx", data_only=True, read_only=True)
    rows = [r for r in wb["Sheet1"].iter_rows(min_row=2, values_only=True) if r[0] is not None]
    wb.close()
    t = np.array([r[0] for r in rows], float)
    R = np.array([r[1] for r in rows], float) * 0.01
    spline = PchipInterpolator(t, R, extrapolate=False)

    def radius(tt):
        scalar = np.ndim(tt) == 0
        tt = np.atleast_1d(np.asarray(tt, float))
        v = spline(np.clip(tt, t[0], t[-1]))
        v = np.where(tt > t[-1], R[-1], v)
        return float(v[0]) if scalar else v

    return radius


def build_grid():
    """材料坐标网格 ξ∈[0,1]：N 个区间、N+1 个节点。"""
    dxi = 1.0 / N
    xi = np.arange(N + 1) * dxi
    faces = np.empty(N + 2)
    faces[1:-1] = 0.5 * (xi[:-1] + xi[1:])
    faces[0], faces[-1] = 0.0, 1.0
    vol = np.pi * (faces[1:] ** 2 - faces[:-1] ** 2)
    area = 2 * np.pi * faces
    return xi, dxi, vol, area[:-1], area[1:]


def harmonic(a):
    s = a[1:] + a[:-1]
    out = np.zeros_like(s)
    nz = s != 0
    out[nz] = 2 * a[1:][nz] * a[:-1][nz] / s[nz]
    return out


def rho_of(c):
    return 760.0 + 90.0 * c


def cp_of(c):
    return 1850.0 + 2150.0 * c / (c + 1.0)


def k_of(c):
    return 0.12 + 0.20 * c / (c + 1.0)


def d_of(c, T):
    return 4.2e-4 * np.exp(-0.30 * c) * np.exp(-3850.0 / (T + 273.15))


def make_rhs(env, radius, geom):
    xi, dxi, vol, al, ar = geom
    n = N + 1

    def fun(t, y):
        R = radius(t)
        T, C = y[:n], y[n:]
        ta, ca = env(t)
        rho, cp, k, d = rho_of(C), cp_of(C), k_of(C), d_of(C, T)
        kh, dh = harmonic(k), harmonic(d)
        qt, qc = np.zeros(n + 1), np.zeros(n + 1)
        qt[1:n] = -kh * (T[1:] - T[:-1]) / dxi
        qc[1:n] = -dh * (C[1:] - C[:-1]) / dxi
        qt[n] = R * HT * (T[-1] - ta)             # 材料坐标下的表面通量
        qc[n] = R * HC * (C[-1] - ca)
        dt = (al * qt[:-1] - ar * qt[1:]) / (R * R * rho * cp * vol)
        dc = (al * qc[:-1] - ar * qc[1:]) / (R * R * vol)
        return np.concatenate((dt, dc))

    return fun


def jac_pattern():
    n = N + 1
    rows, cols = [], []
    for i in range(n):
        for j in (i - 1, i, i + 1):
            if 0 <= j < n:
                rows += [i, i, n + i, n + i]
                cols += [j, n + j, j, n + j]
    return sp.csr_matrix((np.ones(len(rows), bool), (rows, cols)), shape=(2 * n, 2 * n))


class Path:
    def __init__(self, size):
        self.size = size
        self.legs = []

    def add(self, t0, sol):
        self.legs.append((float(t0), float(sol.t[-1]), sol))

    def __call__(self, t):
        t = np.atleast_1d(np.asarray(t, float))
        out = np.full((self.size, t.size), np.nan)
        for t0, t1, sol in self.legs:
            m = (t >= t0 - 1e-9) & (t <= t1 + 1e-9)
            if m.any():
                out[:, m] = sol.sol(t[m])
        return out


def integrate(fun, y0, atol):
    event = lambda t, y: float(np.max(y[N + 1:]) - THRESHOLD)
    event.terminal, event.direction = True, -1.0
    path = Path(2 * (N + 1))
    y, t, hit = y0, 0.0, None
    for t0, t1, step in ((0.0, T_EARLY, MAX_STEP_EARLY), (T_EARLY, T_END, MAX_STEP_LATE)):
        sol = solve_ivp(fun, (max(t0, t), t1), y, method="BDF", rtol=1e-8, atol=atol,
                        jac_sparsity=jac_pattern(), dense_output=True,
                        max_step=step, events=event)
        if not sol.success:
            raise RuntimeError(sol.message)
        path.add(max(t0, t), sol)
        y, t = sol.y[:, -1], float(sol.t[-1])
        if sol.t_events[0].size:
            hit = float(sol.t_events[0][0])
            break
    return path, y, t, hit


def refine(path, t_hit):
    t_lo, t_hi = max(0.0, t_hit - 120.0), t_hit + 150.0
    g = lambda tt: float(np.max(path(np.array([tt]))[N + 1:]) - THRESHOLD)
    return brentq(g, t_lo, t_hi, xtol=1e-12)


def sample(path, xi, radius, times):
    """按 ξ_k=r_k/R(t) 映射回固定物理位置，域外留空；另返回药材表面值。"""
    n = N + 1
    y = path(np.asarray(times, float))
    surf = np.full(times.size, np.nan)
    C = np.full((times.size, len(POS)), np.nan)
    for i, tt in enumerate(times):
        R = float(radius(tt))
        v = y[n:, i]
        inside = POS <= R + 1e-12
        C[i, inside] = np.interp(POS[inside] / R, xi, v)
        surf[i] = v[-1]
    return C, surf


def export(times, data, out_name):
    wb = openpyxl.load_workbook(TEMPLATE_DIR / "result4.xlsx")
    ws = wb["Sheet1"]
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    ws.cell(1, 1).value = "时间\\到药材中心的距离"
    for j, lab in enumerate(LABELS):
        ws.cell(1, 2 + j).value = lab
    for i, tt in enumerate(times):
        cell = ws.cell(2 + i, 1)
        cell.value = float(tt)
        cell.number_format = "0.####"
        for j in range(len(LABELS)):
            v = data[i, j]
            if v is None or not np.isfinite(v):
                continue
            c = ws.cell(2 + i, 2 + j)
            c.value = float(v)
            c.number_format = "0.0000"
    out = ROOT / "result" / out_name
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    wb.close()
    return out


def main():
    env = chamber()
    radius = radius_function()
    geom = build_grid()
    fun = make_rhs(env, radius, geom)
    y0 = np.concatenate((np.full(N + 1, T0), np.full(N + 1, C0)))
    atol = np.concatenate((np.full(N + 1, 1e-8), np.full(N + 1, 1e-10)))

    path, y_end, t_end, t_hit = integrate(fun, y0, atol)
    if t_hit is None:
        raise RuntimeError("在 t=%.0f s 内未触发含水率事件" % T_END)

    extra = solve_ivp(fun, (t_hit, t_hit + 150.0), y_end, method="BDF", rtol=1e-8,
                      atol=atol, jac_sparsity=jac_pattern(), dense_output=True,
                      max_step=MAX_STEP_EARLY)
    path.add(t_hit, extra)
    t_star = refine(path, t_hit)

    times = np.append(np.arange(1, int(np.floor(t_star / 60.0)) + 1) * 60.0, t_star)
    C, surf = sample(path, geom[0], radius, times)
    data = np.column_stack([C, surf])
    out = export(times, data, "result4.xlsx")

    print("问题4  收缩模型（附件2 半径，材料坐标）  N=%d" % N)
    print("  临界烘干时间 t* = %.6f s = %.4f h" % (t_star, t_star / 3600.0))
    print("  临界时刻中心含水率 %.6f，药材表面含水率 %.6f kg/kg" % (C[-1, 0], surf[-1]))
    print("  临界时刻半径 R = %.4f cm（最初 2 cm）" % (radius(t_star) * 100))
    print("  结果已写入 %s" % out)


if __name__ == "__main__":
    main()
