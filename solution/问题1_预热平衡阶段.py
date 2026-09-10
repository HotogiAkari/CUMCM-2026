"""问题1  预热平衡阶段（0~1800 s，固定半径，附录2 常物性）

运行：python solution/问题1_预热平衡阶段.py [N]
输出：result/result1.xlsx（工作表：温度、水分浓度）
"""
from __future__ import annotations

import sys
from pathlib import Path

import numpy as np
import openpyxl
import scipy.sparse as sp
from scipy.integrate import solve_ivp

ROOT = Path(__file__).resolve().parents[1]
INPUT_DIR = ROOT / "inputs" if (ROOT / "inputs").exists() else ROOT
TEMPLATE_DIR = (INPUT_DIR / "templates") if (INPUT_DIR / "templates").exists() else (ROOT / "附件3")

# 半径 m，初始温度 ℃，初始干基含水率 kg/kg，对流换热/传质系数
R0, T0, C0 = 0.02, 28.0, 2.55
HT, HC = 25.0, 8.0e-7
# 附录2 密度 kg/m³、比热容 J/(kg·K)、导热系数 W/(m·K)
RHO, CP, K = 820.0, 2600.0, 0.36

N = int(sys.argv[1]) if len(sys.argv) > 1 else 320
POS = np.arange(21) * 0.001                      # 距中心 0~2 cm，步长 0.1 cm
LABELS = [round(i * 0.1, 10) for i in range(21)]


def chamber():
    """附件1：烘房温度与空气含湿量；4 h 之后取附件末 1 h 均值。"""
    wb = openpyxl.load_workbook(INPUT_DIR / "附件1.xlsx", data_only=True, read_only=True)
    rows = [r for r in wb["Sheet1"].iter_rows(min_row=2, values_only=True) if r[0] is not None]
    wb.close()
    t = np.array([r[0] for r in rows], float)
    ta = np.array([r[1] for r in rows], float)
    ca = np.array([r[2] for r in rows], float)
    tail = t >= t[-1] - 3600.0
    ta_tail, ca_tail = ta[tail].mean(), ca[tail].mean()

    def env(tt):
        scalar = np.ndim(tt) == 0
        tt = np.atleast_1d(np.asarray(tt, float))
        x = np.clip(tt, t[0], t[-1])
        a, b = np.interp(x, t, ta), np.interp(x, t, ca)
        beyond = tt > t[-1]
        a, b = np.where(beyond, ta_tail, a), np.where(beyond, ca_tail, b)
        return (float(a[0]), float(b[0])) if scalar else (a, b)

    return env


def build_grid():
    """节点中心环形控制体：N 个径向区间、N+1 个节点，界面取节点中点。"""
    dr = R0 / N
    r = np.arange(N + 1) * dr
    faces = np.empty(N + 2)
    faces[1:-1] = 0.5 * (r[:-1] + r[1:])
    faces[0], faces[-1] = 0.0, R0
    vol = np.pi * (faces[1:] ** 2 - faces[:-1] ** 2)
    area = 2 * np.pi * faces
    return r, dr, vol, area[:-1], area[1:]


def harmonic(a):
    """界面系数用调和平均。"""
    s = a[1:] + a[:-1]
    out = np.zeros_like(s)
    nz = s != 0
    out[nz] = 2 * a[1:][nz] * a[:-1][nz] / s[nz]
    return out


def make_rhs(env, geom):
    r, dr, vol, al, ar = geom
    n = N + 1

    def fun(t, y):
        T, C = y[:n], y[n:]
        ta, ca = env(t)
        d = 7.0e-9 * np.exp(-0.89 * C)           # 附录2：D(C)
        kh, dh = harmonic(np.full(n, K)), harmonic(d)
        qt, qc = np.zeros(n + 1), np.zeros(n + 1)
        qt[1:n] = -kh * (T[1:] - T[:-1]) / dr     # 内部导热通量
        qc[1:n] = -dh * (C[1:] - C[:-1]) / dr     # 内部扩散通量
        qt[n] = HT * (T[-1] - ta)                 # 表面 Robin 换热
        qc[n] = HC * (C[-1] - ca)                 # 表面 Robin 传质
        dt = (al * qt[:-1] - ar * qt[1:]) / (RHO * CP * vol)
        dc = (al * qc[:-1] - ar * qc[1:]) / vol
        return np.concatenate((dt, dc))

    return fun


def jac_pattern():
    """雅可比稀疏结构：温度块与含水率块均为块三对角。"""
    n = N + 1
    rows, cols = [], []
    for i in range(n):
        for j in (i - 1, i, i + 1):
            if 0 <= j < n:
                rows += [i, i, n + i, n + i]
                cols += [j, n + j, j, n + j]
    return sp.csr_matrix((np.ones(len(rows), bool), (rows, cols)), shape=(2 * n, 2 * n))


def sample(sol, r, times):
    n = N + 1
    y = sol.sol(np.asarray(times, float))
    T = np.array([np.interp(POS, r, v[:n]) for v in y.T])
    C = np.array([np.interp(POS, r, v[n:]) for v in y.T])
    return T, C


def export(sheets, template, out_name):
    """按模板写出结果：数值为 float，显示四位小数。"""
    wb = openpyxl.load_workbook(TEMPLATE_DIR / template)
    for name, (times, data) in sheets.items():
        ws = wb[name]
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row - 1)
        ws.cell(1, 1).value = "时间\\到药材中心的距离"
        for j, lab in enumerate(LABELS):
            ws.cell(1, 2 + j).value = lab
        for i, tt in enumerate(times):
            cell = ws.cell(2 + i, 1)
            cell.value = float(tt)
            cell.number_format = "0.####"
            for j in range(len(LABELS)):
                v = data[i, j]
                if v is None or not np.isfinite(v):
                    continue
                c = ws.cell(2 + i, 2 + j)
                c.value = float(v)
                c.number_format = "0.0000"
    out = ROOT / "result" / out_name
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    wb.close()
    return out


def main():
    env = chamber()
    geom = build_grid()
    y0 = np.concatenate((np.full(N + 1, T0), np.full(N + 1, C0)))
    atol = np.concatenate((np.full(N + 1, 1e-8), np.full(N + 1, 1e-10)))
    sol = solve_ivp(make_rhs(env, geom), (0.0, 1800.0), y0, method="BDF",
                    rtol=1e-8, atol=atol, jac_sparsity=jac_pattern(),
                    dense_output=True, max_step=60.0)
    if not sol.success:
        raise RuntimeError(sol.message)

    times = np.arange(1, 1801, dtype=float)
    T, C = sample(sol, geom[0], times)
    out = export({"温度": (times, T), "水分浓度": (times, C)},
                 "result1.xlsx", "result1.xlsx")

    print("问题1  预热平衡阶段  N=%d" % N)
    print("  t=1800 s  中心温度 %.4f ℃，表面温度 %.4f ℃" % (T[-1, 0], T[-1, -1]))
    print("  t=1800 s  中心含水率 %.4f，表面含水率 %.4f kg/kg" % (C[-1, 0], C[-1, -1]))
    print("  结果已写入 %s" % out)


if __name__ == "__main__":
    main()
