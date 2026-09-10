"""问题2  恒温干燥阶段（0~3 h，固定半径，附录3 变物性）

运行：python solution/问题2_恒温干燥阶段.py [N]
输出：result/result2.xlsx（工作表：温度、水分浓度）
"""
from __future__ import annotations

import sys
from pathlib import Path

import numpy as np
import openpyxl
import scipy.sparse as sp
from scipy.integrate import solve_ivp

ROOT = Path(__file__).resolve().parents[1]
INPUT_DIR = ROOT / "inputs" if (ROOT / "inputs").exists() else ROOT
TEMPLATE_DIR = (INPUT_DIR / "templates") if (INPUT_DIR / "templates").exists() else (ROOT / "附件3")

R0, T0, C0 = 0.02, 28.0, 2.55
HT, HC = 25.0, 8.0e-7

N = int(sys.argv[1]) if len(sys.argv) > 1 else 320
POS = np.arange(21) * 0.001
LABELS = [round(i * 0.1, 10) for i in range(21)]


def chamber():
    """附件1：烘房温度与空气含湿量；4 h 之后取附件末 1 h 均值。"""
    wb = openpyxl.load_workbook(INPUT_DIR / "附件1.xlsx", data_only=True, read_only=True)
    rows = [r for r in wb["Sheet1"].iter_rows(min_row=2, values_only=True) if r[0] is not None]
    wb.close()
    t = np.array([r[0] for r in rows], float)
    ta = np.array([r[1] for r in rows], float)
    ca = np.array([r[2] for r in rows], float)
    tail = t >= t[-1] - 3600.0
    ta_tail, ca_tail = ta[tail].mean(), ca[tail].mean()

    def env(tt):
        scalar = np.ndim(tt) == 0
        tt = np.atleast_1d(np.asarray(tt, float))
        x = np.clip(tt, t[0], t[-1])
        a, b = np.interp(x, t, ta), np.interp(x, t, ca)
        beyond = tt > t[-1]
        a, b = np.where(beyond, ta_tail, a), np.where(beyond, ca_tail, b)
        return (float(a[0]), float(b[0])) if scalar else (a, b)

    return env


def build_grid():
    dr = R0 / N
    r = np.arange(N + 1) * dr
    faces = np.empty(N + 2)
    faces[1:-1] = 0.5 * (r[:-1] + r[1:])
    faces[0], faces[-1] = 0.0, R0
    vol = np.pi * (faces[1:] ** 2 - faces[:-1] ** 2)
    area = 2 * np.pi * faces
    return r, dr, vol, area[:-1], area[1:]


def harmonic(a):
    s = a[1:] + a[:-1]
    out = np.zeros_like(s)
    nz = s != 0
    out[nz] = 2 * a[1:][nz] * a[:-1][nz] / s[nz]
    return out


# 附录3 物性：密度、比热容、导热系数随含水率变化
def rho_of(C):
    return 650.0 + 128.0 * C


def cp_of(C):
    return 1450.0 + 2736.0 * C / (C + 1.0)


def k_of(C):
    return 0.21 + 0.38 * C / (C + 1.0)


def d_of(C, T):
    """附录3 扩散系数，温度项取绝对温度倒数形式。"""
    return 2.4e-3 * np.exp(-0.45 * C) * np.exp(-3850.0 / (T + 273.15))


def make_rhs(env, geom):
    r, dr, vol, al, ar = geom
    n = N + 1

    def fun(t, y):
        T, C = y[:n], y[n:]
        ta, ca = env(t)
        rho, cp, k, d = rho_of(C), cp_of(C), k_of(C), d_of(C, T)
        kh, dh = harmonic(k), harmonic(d)
        qt, qc = np.zeros(n + 1), np.zeros(n + 1)
        qt[1:n] = -kh * (T[1:] - T[:-1]) / dr
        qc[1:n] = -dh * (C[1:] - C[:-1]) / dr
        qt[n] = HT * (T[-1] - ta)
        qc[n] = HC * (C[-1] - ca)
        dt = (al * qt[:-1] - ar * qt[1:]) / (rho * cp * vol)
        dc = (al * qc[:-1] - ar * qc[1:]) / vol
        return np.concatenate((dt, dc))

    return fun


def jac_pattern():
    n = N + 1
    rows, cols = [], []
    for i in range(n):
        for j in (i - 1, i, i + 1):
            if 0 <= j < n:
                rows += [i, i, n + i, n + i]
                cols += [j, n + j, j, n + j]
    return sp.csr_matrix((np.ones(len(rows), bool), (rows, cols)), shape=(2 * n, 2 * n))


def sample(sol, r, times):
    n = N + 1
    y = sol.sol(np.asarray(times, float))
    T = np.array([np.interp(POS, r, v[:n]) for v in y.T])
    C = np.array([np.interp(POS, r, v[n:]) for v in y.T])
    return T, C


def export(sheets, template, out_name):
    wb = openpyxl.load_workbook(TEMPLATE_DIR / template)
    for name, (times, data) in sheets.items():
        ws = wb[name]
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row - 1)
        ws.cell(1, 1).value = "时间\\到药材中心的距离"
        for j, lab in enumerate(LABELS):
            ws.cell(1, 2 + j).value = lab
        for i, tt in enumerate(times):
            cell = ws.cell(2 + i, 1)
            cell.value = float(tt)
            cell.number_format = "0.####"
            for j in range(len(LABELS)):
                v = data[i, j]
                if v is None or not np.isfinite(v):
                    continue
                c = ws.cell(2 + i, 2 + j)
                c.value = float(v)
                c.number_format = "0.0000"
    out = ROOT / "result" / out_name
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    wb.close()
    return out


def main():
    env = chamber()
    geom = build_grid()
    y0 = np.concatenate((np.full(N + 1, T0), np.full(N + 1, C0)))
    atol = np.concatenate((np.full(N + 1, 1e-8), np.full(N + 1, 1e-10)))
    sol = solve_ivp(make_rhs(env, geom), (0.0, 10800.0), y0, method="BDF",
                    rtol=1e-8, atol=atol, jac_sparsity=jac_pattern(),
                    dense_output=True, max_step=60.0)
    if not sol.success:
        raise RuntimeError(sol.message)

    times = np.arange(1, 10801, dtype=float)
    T, C = sample(sol, geom[0], times)
    out = export({"温度": (times, T), "水分浓度": (times, C)},
                 "result2.xlsx", "result2.xlsx")

    print("问题2  恒温干燥阶段  N=%d" % N)
    print("  t=3 h  中心温度 %.4f ℃，表面温度 %.4f ℃" % (T[-1, 0], T[-1, -1]))
    print("  t=3 h  中心含水率 %.4f，表面含水率 %.4f kg/kg" % (C[-1, 0], C[-1, -1]))
    print("  结果已写入 %s" % out)


if __name__ == "__main__":
    main()
