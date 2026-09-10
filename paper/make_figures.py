#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成论文插图：读取 result*.xlsx 与诊断 CSV，输出 figures/*.png。

若本机无 MATLAB/Octave，先用本脚本出图；matlab/plot_results.m 为等价 MATLAB 版本。
"""
from __future__ import annotations

import csv
import warnings
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
FIG = ROOT / "figures"
FIG.mkdir(exist_ok=True)

plt.rcParams.update({
    "font.sans-serif": ["Noto Sans CJK JP", "DejaVu Sans"],
    "axes.unicode_minus": False,
    "font.size": 10,
    "figure.dpi": 150,
    "savefig.bbox": "tight",
})

T_STAR3, T_STAR4 = 59286.175088758, 70759.978749528


def sheet(path, name):
    wb = openpyxl.load_workbook(ROOT / "result" / path, data_only=True, read_only=True)
    ws = wb[name]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    cols = rows[0][1:]
    ncol = len(cols)
    r = np.array([v for v in cols if isinstance(v, (int, float))], float) / 100.0
    times = np.array([row[0] for row in rows[1:]], float)
    data = np.full((len(rows) - 1, ncol), np.nan)
    for i, row in enumerate(rows[1:]):
        vals = list(row[1:1 + ncol])
        for j, v in enumerate(vals):
            if isinstance(v, (int, float)):
                data[i, j] = v
    return r, times, data


def row_at(times, data, t):
    return data[int(np.argmin(np.abs(times - t)))]


def csv_rows(path):
    with open(ROOT / "outputs" / "diagnostics" / path, newline="", encoding="utf-8") as fh:
        return list(csv.reader(fh))


def fig1():
    r, times, T = sheet("result1.xlsx", "温度")
    _, _, C = sheet("result1.xlsx", "水分浓度")
    fig, ax = plt.subplots(1, 2, figsize=(9, 3.2))
    for t in (100, 600, 1200, 1800):
        ax[0].plot(r, row_at(times, T, t), marker="o", ms=3, label="t=%d s" % t)
        ax[1].plot(r, row_at(times, C, t), marker="s", ms=3, label="t=%d s" % t)
    ax[0].set_xlabel("到中心距离 r/cm"), ax[0].set_ylabel("温度 T/℃")
    ax[1].set_xlabel("到中心距离 r/cm"), ax[1].set_ylabel("干基含水率 C/(kg/kg)")
    for a in ax:
        a.grid(alpha=0.3), a.legend(fontsize=8)
    fig.suptitle("图1  问题一：预热平衡阶段温度与含水率的径向分布", fontsize=11)
    fig.tight_layout(), fig.savefig(FIG / "fig1_q1_radial.png"), plt.close(fig)


def fig2():
    r, times, T = sheet("result2.xlsx", "温度")
    _, _, C = sheet("result2.xlsx", "水分浓度")
    fig, ax = plt.subplots(1, 2, figsize=(9, 3.2))
    for h in (0.5, 1.0, 1.5, 2.0, 3.0):
        ax[0].plot(r, row_at(times, T, h * 3600), marker="o", ms=3, label="t=%g h" % h)
        ax[1].plot(r, row_at(times, C, h * 3600), marker="s", ms=3, label="t=%g h" % h)
    ax[0].set_xlabel("到中心距离 r/cm"), ax[0].set_ylabel("温度 T/℃")
    ax[1].set_xlabel("到中心距离 r/cm"), ax[1].set_ylabel("干基含水率 C/(kg/kg)")
    for a in ax:
        a.grid(alpha=0.3), a.legend(fontsize=8)
    fig.suptitle("图2  问题二：恒温干燥阶段温度与含水率的径向分布", fontsize=11)
    fig.tight_layout(), fig.savefig(FIG / "fig2_q2_radial.png"), plt.close(fig)


def fig3():
    r, times, C = sheet("result3.xlsx", "Sheet1")
    th = times / 3600.0
    centre, surface = C[:, 0], C[:, len(r) - 1]
    fig, ax = plt.subplots(figsize=(7, 3.6))
    ax.plot(th, centre, label="中心 r=0")
    ax.plot(th, surface, label="表面 r=2 cm")
    ax.axhline(0.15, ls="--", c="crimson", label="判据 max C = 0.15")
    ax.axvline(T_STAR3 / 3600.0, ls=":", c="gray")
    ax.annotate("t*=%.4f h" % (T_STAR3 / 3600),
                xy=(T_STAR3 / 3600, 0.15), xytext=(T_STAR3 / 3600 - 5.5, 0.85),
                arrowprops=dict(arrowstyle="->", color="gray"), fontsize=9)
    ax.set_xlabel("时间 t/h"), ax.set_ylabel("干基含水率 C/(kg/kg)")
    ax.set_title("图3  问题三：含水率演化与烘干终止判据"), ax.grid(alpha=0.3), ax.legend(fontsize=9)
    fig.tight_layout(), fig.savefig(FIG / "fig3_q3_history.png"), plt.close(fig)


def fig4():
    r, times, C = sheet("result4.xlsx", "Sheet1")
    th = times / 3600.0
    centre = C[:, 0]
    surface = C[:, -1]
    wb = openpyxl.load_workbook(ROOT / "inputs" / "附件2.xlsx", data_only=True, read_only=True)
    rows = [r_ for r_ in wb["Sheet1"].iter_rows(min_row=2, values_only=True) if r_[0] is not None]
    wb.close()
    rt = np.array([x[0] for x in rows], float) / 3600.0
    rr = np.array([x[1] for x in rows], float)

    fig, ax = plt.subplots(figsize=(7, 3.6))
    ax.plot(th, centre, label="中心 r=0")
    ax.plot(th, surface, "C1", label="药材表面")
    ax.axhline(0.15, ls="--", c="crimson", label="判据 max C = 0.15")
    ax.axvline(T_STAR4 / 3600.0, ls=":", c="gray")
    ax.annotate("t*=%.4f h" % (T_STAR4 / 3600),
                xy=(T_STAR4 / 3600, 0.15), xytext=(T_STAR4 / 3600 - 7.5, 0.9),
                arrowprops=dict(arrowstyle="->", color="gray"), fontsize=9)
    ax.set_xlabel("时间 t/h"), ax.set_ylabel("干基含水率 C/(kg/kg)")
    ax.grid(alpha=0.3)
    ax2 = ax.twinx()
    ax2.plot(rt, rr, "g--", alpha=0.7, label="半径 R(t)")
    ax2.set_ylabel("药材半径 R/cm", color="g")
    ax2.tick_params(axis="y", colors="g")
    h1, l1 = ax.get_legend_handles_labels()
    h2, l2 = ax2.get_legend_handles_labels()
    ax.legend(h1 + h2, l1 + l2, fontsize=9, loc="lower left")
    ax.set_title("图4  问题四：收缩条件下的含水率演化与半径变化")
    fig.tight_layout(), fig.savefig(FIG / "fig4_q4_history.png"), plt.close(fig)


def fig5():
    mt = csv_rows("sensitivity_mass_transfer.csv")[1:]
    df = csv_rows("sensitivity_diffusivity.csv")[1:]
    ca = csv_rows("sensitivity_equilibrium_moisture.csv")[1:]
    fig, ax = plt.subplots(1, 3, figsize=(11, 3.2))
    sets = [("传质系数 $h_C$ 倍数", mt, 0), ("扩散系数 $D$ 倍数", df, 0),
            ("平衡含水率 $C_a$ 倍数", ca, 0)]
    for a, (title, rows, _) in zip(ax, sets):
        f = [float(r[0]) for r in rows]
        if rows is mt:
            t3 = [float(r[3]) for r in rows]
            t4 = [float(r[5]) for r in rows]
        else:
            t3 = [float(r[2]) for r in rows]
            t4 = [float(r[4]) for r in rows]
        a.plot(f, t3, "o-", label="问题三")
        a.plot(f, t4, "s-", label="问题四")
        a.set_xlabel(title), a.set_ylabel("临界时间 $t^*$/h")
        a.grid(alpha=0.3), a.legend(fontsize=8)
    fig.suptitle("图5  关键参数灵敏度：临界烘干时间随参数倍数的变化", fontsize=11)
    fig.tight_layout(), fig.savefig(FIG / "fig5_sensitivity.png"), plt.close(fig)


def fig6():
    sp = csv_rows("space.csv")[1:]
    tol = csv_rows("tolerance.csv")[1:]
    fig, ax = plt.subplots(1, 2, figsize=(9.5, 3.2))
    for q, mk in (("Q3", "o-"), ("Q4", "s-")):
        rows = [r for r in sp if r[0] == q]
        ax[0].plot([int(r[1]) for r in rows], [float(r[3]) for r in rows], mk, label=q)
    ax[0].set_xlabel("径向区间数 N"), ax[0].set_ylabel("临界时间 $t^*$/s")
    ax[0].set_title("(a) 空间网格收敛"), ax[0].grid(alpha=0.3), ax[0].legend(fontsize=9)
    ax[1].plot([float(r[0]) for r in tol], [float(r[2]) for r in tol], "o-", label="问题三")
    ax[1].plot([float(r[0]) for r in tol], [float(r[4]) for r in tol], "s-", label="问题四")
    ax[1].set_xscale("log"), ax[1].set_xlabel("相对容差 rtol"), ax[1].set_ylabel("临界时间 $t^*$/h")
    ax[1].set_title("(b) 时间容差收敛"), ax[1].grid(alpha=0.3), ax[1].legend(fontsize=9)
    fig.suptitle("图6  数值收敛性检验", fontsize=11)
    fig.tight_layout(), fig.savefig(FIG / "fig6_convergence.png"), plt.close(fig)


if __name__ == "__main__":
    with warnings.catch_warnings(record=True) as w:
        warnings.simplefilter("always")
        for fn in (fig1, fig2, fig3, fig4, fig5, fig6):
            fn()
        missing = [str(x.message) for x in w if "missing from font" in str(x.message).lower()
                   or "Glyph" in str(x.message)]
    for p in sorted(FIG.glob("*.png")):
        print("wrote", p, "%.0f KB" % (p.stat().st_size / 1024))
    print("缺字告警:", len(missing))
    for m in missing[:5]:
        print("  ", m)
