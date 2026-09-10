"""Write results into copies of the official Excel templates."""
from __future__ import annotations

from pathlib import Path

import numpy as np
import openpyxl

HEADER_LABEL = "时间\\到药材中心的距离"


def _clear_data_rows(ws):
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)


def export_result(template_path, out_path, sheet_payloads, decimals: int = 4):
    """Copy the template and fill it with the given payloads.

    sheet_payloads: {sheet_name: {"times": array, "labels": [col labels],
                                  "data": (n_time, n_lab) array with NaN = blank}}
    """
    template_path = Path(template_path)
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.load_workbook(template_path)
    try:
        for sheet_name, payload in sheet_payloads.items():
            ws = wb[sheet_name]
            _clear_data_rows(ws)
            times = np.asarray(payload["times"], float)
            labels = list(payload["labels"])
            data = np.asarray(payload["data"], float)

            ws.cell(1, 1).value = HEADER_LABEL
            for j, lab in enumerate(labels):
                ws.cell(1, 2 + j).value = lab

            fmt = "0." + "0" * decimals
            for i in range(times.size):
                tc = ws.cell(2 + i, 1)
                tc.value = float(times[i])
                tc.number_format = "0.####"
                for j in range(len(labels)):
                    v = data[i, j]
                    if v is None or not np.isfinite(v):
                        continue
                    cell = ws.cell(2 + i, 2 + j)
                    cell.value = float(v)
                    cell.number_format = fmt
        wb.save(out_path)
    finally:
        wb.close()
    return out_path
