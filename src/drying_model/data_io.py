"""Read and validate the raw attachments and the Excel result templates.

Raw attachments are always opened read-only.
"""
from __future__ import annotations

import numpy as np
import openpyxl


def _load_sheet_rows(path, sheet: str, ncol: int) -> np.ndarray:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        if sheet not in wb.sheetnames:
            raise KeyError("sheet %r not found in %s (have %s)" % (sheet, path, wb.sheetnames))
        ws = wb[sheet]
        rows = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row is None or row[0] is None:
                continue
            rows.append([row[i] for i in range(ncol)])
    finally:
        wb.close()
    return np.asarray(rows, dtype=float)


def read_attachment1(path):
    """Return (time_s, chamber_T_C, chamber_moisture)."""
    arr = _load_sheet_rows(path, "Sheet1", 3)
    t, ta, ca = arr[:, 0], arr[:, 1], arr[:, 2]
    if t.size != 241:
        raise ValueError("attachment1: expected 241 records, found %d" % t.size)
    if t[0] != 0.0 or t[-1] != 14400.0:
        raise ValueError("attachment1: time span must be 0..14400 s, found %.1f..%.1f" % (t[0], t[-1]))
    if not np.all(np.diff(t) > 0):
        raise ValueError("attachment1: time must be strictly increasing")
    if not np.isfinite(arr).all():
        raise ValueError("attachment1: contains NaN or Inf")
    return t, ta, ca


def read_attachment2(path):
    """Return (time_s, radius_cm)."""
    arr = _load_sheet_rows(path, "Sheet1", 2)
    t, r_cm = arr[:, 0], arr[:, 1]
    if t.size != 145:
        raise ValueError("attachment2: expected 145 records, found %d" % t.size)
    if not np.all(np.diff(t) > 0):
        raise ValueError("attachment2: time must be strictly increasing")
    if not np.isfinite(arr).all():
        raise ValueError("attachment2: contains NaN or Inf")
    if not np.all(r_cm > 0):
        raise ValueError("attachment2: radius must be positive")
    if not np.all(np.diff(r_cm) <= 1e-12):
        raise ValueError("attachment2: radius must be non-increasing")
    return t, r_cm


def read_template_info(path) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True)
    info = {"path": str(path), "sheets": []}
    try:
        for ws in wb.worksheets:
            hdr = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
            times = [ws.cell(r, 1).value for r in range(2, ws.max_row + 1)]
            info["sheets"].append({
                "title": ws.title,
                "max_row": ws.max_row,
                "max_col": ws.max_column,
                "header": hdr,
                "times": times,
                "merged": [str(m) for m in ws.merged_cells.ranges],
            })
    finally:
        wb.close()
    return info
